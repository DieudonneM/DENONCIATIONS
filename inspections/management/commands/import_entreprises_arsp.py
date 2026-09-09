from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from openpyxl import load_workbook

from inspections.models import Entreprise


REQUIRED_COLUMNS = {'Dénomination', 'RCCM', 'Num Impôt'}


def normalize(value):
    """Convertit les cellules Excel en texte comparable."""
    return ' '.join(str(value or '').strip().split())


class Command(BaseCommand):
    help = 'Importe les entreprises ARSP (Dénomination, RCCM et Num Impôt uniquement).'

    def add_arguments(self, parser):
        parser.add_argument(
            '--file',
            default='inspections/liste_entreprises_arsp.xlsx',
            help='Chemin du fichier Excel ARSP, relatif au projet ou absolu.',
        )
        parser.add_argument(
            '--dry-run',
            action='store_true',
            help='Affiche le bilan sans enregistrer de donnees.',
        )

    def handle(self, *args, **options):
        source_path = Path(options['file'])
        if not source_path.is_absolute():
            from django.conf import settings

            source_path = settings.BASE_DIR / source_path
        if not source_path.is_file():
            raise CommandError(f'Fichier Excel introuvable : {source_path}')

        workbook = load_workbook(source_path, read_only=True, data_only=True)
        worksheet = workbook.active
        headers = [normalize(cell) for cell in next(worksheet.iter_rows(values_only=True))]
        missing_columns = REQUIRED_COLUMNS - set(headers)
        if missing_columns:
            raise CommandError(
                'Colonnes obligatoires absentes : ' + ', '.join(sorted(missing_columns))
            )
        column_indexes = {header: index for index, header in enumerate(headers)}

        created = updated = skipped = 0
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            nom = normalize(row[column_indexes['Dénomination']])
            rccm = normalize(row[column_indexes['RCCM']])
            nif = normalize(row[column_indexes['Num Impôt']])
            if not nom:
                skipped += 1
                continue

            entreprise = None
            if rccm:
                entreprise = Entreprise.objects.filter(rccm__iexact=rccm).first()
            if entreprise is None and nif:
                entreprise = Entreprise.objects.filter(nif__iexact=nif).first()
            if entreprise is None:
                entreprise = Entreprise.objects.filter(nom__iexact=nom).first()

            if entreprise is None:
                created += 1
                if not options['dry_run']:
                    Entreprise.objects.create(nom=nom, rccm=rccm, nif=nif)
                continue

            changed_fields = []
            if not entreprise.rccm and rccm:
                entreprise.rccm = rccm
                changed_fields.append('rccm')
            if not entreprise.nif and nif:
                entreprise.nif = nif
                changed_fields.append('nif')
            if changed_fields:
                updated += 1
                if not options['dry_run']:
                    entreprise.save(update_fields=changed_fields)
            else:
                skipped += 1

        workbook.close()
        action = 'Simulation terminee' if options['dry_run'] else 'Importation terminee'
        self.stdout.write(self.style.SUCCESS(
            f'{action} : {created} creee(s), {updated} mise(s) a jour, {skipped} ignoree(s).'
        ))